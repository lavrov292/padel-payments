from __future__ import annotations

import hashlib
import base64
from io import BytesIO
import json
import os
import re
import secrets
import sqlite3
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from fastapi import FastAPI, HTTPException, Query
from fastapi.responses import HTMLResponse, Response


DB_ENV = "LUNDA_DB_PATH"
LEVEL_ORDER = ["D", "D+", "C", "C+", "B", "B+", "A"]
LEVEL_INDEX = {level: idx for idx, level in enumerate(LEVEL_ORDER)}
TOURNAMENT_TYPES = [
    {"value": "individual", "label": "Индивидуальный"},
    {"value": "pair", "label": "Парный"},
    {"value": "women", "label": "Женский"},
]
TIME_PERIODS = [
    {"value": "morning", "label": "Утро"},
    {"value": "day", "label": "День"},
    {"value": "evening", "label": "Вечер"},
]


app = FastAPI(title="Lunda Stats", version="0.1.0")


@app.middleware("http")
async def optional_basic_auth(request, call_next):
    user = os.environ.get("LUNDA_WEB_USER")
    password = os.environ.get("LUNDA_WEB_PASSWORD")
    if not user and not password:
        response = await call_next(request)
        response.headers["Cache-Control"] = "no-store"
        return response

    header = request.headers.get("authorization", "")
    expected = f"{user}:{password}"
    authorized = False
    if header.lower().startswith("basic "):
        try:
            decoded = base64.b64decode(header.split(" ", 1)[1]).decode("utf-8")
            authorized = secrets.compare_digest(decoded, expected)
        except Exception:
            authorized = False

    if not authorized:
        return Response(
            "Authentication required",
            status_code=401,
            headers={"WWW-Authenticate": 'Basic realm="Lunda Stats"'},
        )

    response = await call_next(request)
    response.headers["Cache-Control"] = "no-store"
    return response


def db_path() -> Path:
    configured = os.environ.get(DB_ENV)
    if configured:
        return Path(configured).expanduser()

    return Path.cwd() / "outputs" / "participant_window_2026-07-08" / "participant_window_snapshot.sqlite3"


def connect() -> sqlite3.Connection:
    path = db_path()
    if not path.exists():
        raise HTTPException(status_code=500, detail=f"SQLite database not found: {path}")
    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    return conn


def split_multi(values: list[str] | None) -> list[str]:
    if not values:
        return []
    result: list[str] = []
    for value in values:
        result.extend(part.strip() for part in value.split(",") if part.strip())
    return result


def parse_skill_levels(value: str | None) -> set[str]:
    if not value:
        return set()

    normalized = value.upper().replace("…", "..").replace("—", "-").replace("–", "-")
    found = re.findall(r"\b[DCBA]\+?\b", normalized)
    if not found:
        return set()

    if len(found) >= 2 and (".." in normalized or "-" in normalized):
        start = LEVEL_INDEX.get(found[0])
        end = LEVEL_INDEX.get(found[-1])
        if start is not None and end is not None:
            lo, hi = sorted((start, end))
            return set(LEVEL_ORDER[lo : hi + 1])

    return {level for level in found if level in LEVEL_INDEX}


def level_matches(skill_level: str | None, selected_levels: set[str]) -> bool:
    if not selected_levels:
        return True
    return bool(parse_skill_levels(skill_level) & selected_levels)


def tournament_type(row: sqlite3.Row | dict[str, Any]) -> str:
    title = str(row["title"] or "").strip().lower()
    unit = str(row["participants_unit"] or "").strip().lower()
    if title.startswith("женский турнир"):
        return "women"
    if unit.startswith("команд"):
        return "pair"
    return "individual"


def tournament_type_label(value: str) -> str:
    labels = {item["value"]: item["label"] for item in TOURNAMENT_TYPES}
    return labels.get(value, value)


def time_period(row: sqlite3.Row | dict[str, Any]) -> str:
    starts_at = row["starts_at"]
    if not starts_at:
        return ""
    try:
        start = datetime.fromisoformat(str(starts_at))
    except ValueError:
        return ""
    minutes = start.hour * 60 + start.minute
    if minutes < 11 * 60:
        return "morning"
    if 12 * 60 <= minutes < 17 * 60:
        return "day"
    if minutes >= 17 * 60:
        return "evening"
    return ""


def time_period_label(value: str) -> str:
    labels = {item["value"]: item["label"] for item in TIME_PERIODS}
    return labels.get(value, value)


def tournament_matches_extra_filters(
    row: sqlite3.Row,
    *,
    selected_levels: set[str],
    selected_types: set[str],
    selected_periods: set[str],
) -> bool:
    if not level_matches(row["skill_level"], selected_levels):
        return False
    if selected_types and tournament_type(row) not in selected_types:
        return False
    if selected_periods and time_period(row) not in selected_periods:
        return False
    return True


def parse_iso_date(value: str | None) -> date | None:
    if not value:
        return None
    return date.fromisoformat(value)


def today_iso() -> str:
    return date.today().isoformat()


def week_bounds(anchor: str | None = None) -> tuple[str, str]:
    current = parse_iso_date(anchor) or date.today()
    start = current - timedelta(days=current.weekday())
    end = start + timedelta(days=6)
    return start.isoformat(), end.isoformat()


def append_in_filter(filters: list[str], params: list[Any], column: str, values: list[str]) -> None:
    if not values:
        return
    placeholders = ",".join("?" for _ in values)
    filters.append(f"{column} IN ({placeholders})")
    params.extend(values)


def tournament_filters(
    *,
    date_from: str = "",
    date_to: str = "",
    locations: list[str] | None = None,
    organizers: list[str] | None = None,
    formats: list[str] | None = None,
    include_cancelled: bool = False,
) -> tuple[str, list[Any]]:
    filters: list[str] = []
    params: list[Any] = []
    if not include_cancelled:
        filters.append("t.source_status NOT IN ('cancelled', 'missing')")
    if date_from:
        filters.append("t.tournament_date >= ?")
        params.append(date_from)
    if date_to:
        filters.append("t.tournament_date <= ?")
        params.append(date_to)
    append_in_filter(filters, params, "t.location", locations or [])
    append_in_filter(filters, params, "t.organizer", organizers or [])
    append_in_filter(filters, params, "t.format", formats or [])
    where_sql = f"WHERE {' AND '.join(filters)}" if filters else ""
    return where_sql, params


def organizer_color(organizer: str | None) -> str:
    palette = [
        "#2563eb",
        "#059669",
        "#dc2626",
        "#7c3aed",
        "#ca8a04",
        "#0891b2",
        "#db2777",
        "#4f46e5",
        "#16a34a",
        "#ea580c",
    ]
    key = organizer or ""
    digest = hashlib.sha1(key.encode("utf-8")).hexdigest()
    return palette[int(digest[:2], 16) % len(palette)]


def row_to_tournament(row: sqlite3.Row) -> dict[str, Any]:
    result = dict(row)
    type_value = tournament_type(row)
    period_value = time_period(row)
    result.pop("raw_json", None)
    result["level_tags"] = sorted(parse_skill_levels(result.get("skill_level")), key=lambda item: LEVEL_INDEX[item])
    result["tournament_type"] = type_value
    result["tournament_type_label"] = tournament_type_label(type_value)
    result["time_period"] = period_value
    result["time_period_label"] = time_period_label(period_value)
    result["color"] = organizer_color(result.get("organizer"))
    result["starts_at_display"] = format_time(result.get("starts_at"))
    result["ends_at"] = infer_ends_at(result.get("starts_at"), result.get("time_label"))
    result["ends_at_display"] = format_time(result.get("ends_at"))
    return result


def is_valid_player_name(value: str) -> bool:
    text = value.strip()
    if not text:
        return False
    letters = re.sub(r"[^A-Za-zА-Яа-яЁё]", "", text)
    if len(letters) < 2:
        return False
    if text in {"-", "—", "–"}:
        return False
    return True


def format_time(value: str | None) -> str:
    if not value:
        return ""
    try:
        return datetime.fromisoformat(value).strftime("%H:%M")
    except ValueError:
        return ""


def infer_ends_at(starts_at: str | None, time_label: str | None) -> str | None:
    if not starts_at:
        return None
    try:
        start = datetime.fromisoformat(starts_at)
    except ValueError:
        return None

    if time_label:
        matches = re.findall(r"(\d{1,2})[:.](\d{2})", time_label)
        if len(matches) >= 2:
            hour, minute = map(int, matches[-1])
            end = start.replace(hour=hour, minute=minute)
            if end <= start:
                end += timedelta(days=1)
            return end.isoformat()

    return (start + timedelta(hours=2)).isoformat()


@app.get("/", response_class=HTMLResponse)
def index() -> str:
    return HTML


@app.get("/api/filters")
def filters() -> dict[str, Any]:
    conn = connect()
    try:
        def distinct(column: str) -> list[str]:
            rows = conn.execute(
                f"""
                SELECT DISTINCT {column} AS value
                FROM tournaments
                WHERE {column} IS NOT NULL AND TRIM({column}) != ''
                ORDER BY value
                """
            ).fetchall()
            return [str(row["value"]) for row in rows]

        return {
            "locations": distinct("location"),
            "organizers": distinct("organizer"),
            "formats": distinct("format"),
            "levels": LEVEL_ORDER,
            "tournament_types": TOURNAMENT_TYPES,
            "time_periods": TIME_PERIODS,
            "date_min": conn.execute("SELECT MIN(tournament_date) FROM tournaments").fetchone()[0],
            "date_max": conn.execute("SELECT MAX(tournament_date) FROM tournaments").fetchone()[0],
        }
    finally:
        conn.close()


@app.get("/api/players")
def players(
    date_from: str = "",
    date_to: str = "",
    location: list[str] | None = Query(default=None),
    organizer: list[str] | None = Query(default=None),
    level: list[str] | None = Query(default=None),
    format_value: list[str] | None = Query(default=None, alias="format"),
    tournament_type_value: list[str] | None = Query(default=None, alias="tournament_type"),
    time_period_value: list[str] | None = Query(default=None, alias="time_period"),
    rating_min: float | None = Query(default=None, ge=1, le=7),
    rating_max: float | None = Query(default=None, ge=1, le=7),
    search: str = "",
    limit: int = Query(default=500, ge=1, le=5000),
) -> dict[str, Any]:
    selected_levels = set(split_multi(level))
    selected_types = set(split_multi(tournament_type_value))
    selected_periods = set(split_multi(time_period_value))
    where_sql, params = tournament_filters(
        date_from=date_from,
        date_to=date_to,
        locations=split_multi(location),
        organizers=split_multi(organizer),
        formats=split_multi(format_value),
    )
    conn = connect()
    try:
        rows = conn.execute(
            f"""
            SELECT
                COALESCE(p.id, 0) AS player_id,
                COALESCE(p.display_name, fp.raw_name) AS player_name,
                fp.normalized_name,
                fp.rating,
                p.latest_rating,
                t.id AS tournament_id,
                t.tournament_date,
                t.time_label,
                t.title,
                t.organizer,
                t.location,
                t.skill_level,
                t.format,
                t.participants_unit,
                t.starts_at
            FROM final_participations fp
            JOIN tournaments t ON t.id = fp.tournament_id
            LEFT JOIN players p ON p.id = fp.player_id
            {where_sql}
            ORDER BY player_name, t.starts_at
            """,
            params,
        ).fetchall()

        search_lower = search.strip().lower()
        grouped: dict[str, dict[str, Any]] = {}
        for row in rows:
            if not tournament_matches_extra_filters(
                row,
                selected_levels=selected_levels,
                selected_types=selected_types,
                selected_periods=selected_periods,
            ):
                continue
            player_name = str(row["player_name"] or "")
            if not is_valid_player_name(player_name):
                continue
            if search_lower and search_lower not in player_name.lower():
                continue
            key = f"{row['player_id']}|{row['normalized_name']}"
            item = grouped.setdefault(
                key,
                {
                    "player_name": player_name,
                    "tournament_count": 0,
                    "locations": set(),
                    "organizers": set(),
                    "levels": set(),
                    "formats": set(),
                    "latest_rating": None,
                    "latest_rating_at": "",
                    "tournaments": [],
                },
            )
            item["tournament_count"] += 1
            if row["rating"] is not None and str(row["starts_at"] or "") >= str(item["latest_rating_at"] or ""):
                item["latest_rating"] = row["rating"]
                item["latest_rating_at"] = row["starts_at"] or ""
            elif item["latest_rating"] is None and row["latest_rating"] is not None:
                item["latest_rating"] = row["latest_rating"]
            for field, target in [
                ("location", "locations"),
                ("organizer", "organizers"),
                ("skill_level", "levels"),
                ("format", "formats"),
            ]:
                if row[field]:
                    item[target].add(row[field])
            item["tournaments"].append(
                {
                    "id": row["tournament_id"],
                    "date": row["tournament_date"],
                    "time": row["time_label"],
                    "title": row["title"],
                    "organizer": row["organizer"],
                    "location": row["location"],
                    "level": row["skill_level"],
                    "format": row["format"],
                }
            )

        result = []
        for item in grouped.values():
            latest_rating = item["latest_rating"]
            if rating_min is not None and (latest_rating is None or float(latest_rating) < rating_min):
                continue
            if rating_max is not None and (latest_rating is None or float(latest_rating) > rating_max):
                continue
            result.append(
                {
                    **item,
                    "latest_rating": latest_rating,
                    "locations": sorted(item["locations"]),
                    "organizers": sorted(item["organizers"]),
                    "levels": sorted(item["levels"]),
                    "formats": sorted(item["formats"]),
                }
            )

        result.sort(key=lambda item: (-item["tournament_count"], item["player_name"].lower()))
        return {"items": result[:limit], "total": len(result)}
    finally:
        conn.close()


@app.get("/api/players/export")
def export_players(
    date_from: str = "",
    date_to: str = "",
    location: list[str] | None = Query(default=None),
    organizer: list[str] | None = Query(default=None),
    level: list[str] | None = Query(default=None),
    format_value: list[str] | None = Query(default=None, alias="format"),
    tournament_type_value: list[str] | None = Query(default=None, alias="tournament_type"),
    time_period_value: list[str] | None = Query(default=None, alias="time_period"),
    rating_min: float | None = Query(default=None, ge=1, le=7),
    rating_max: float | None = Query(default=None, ge=1, le=7),
    search: str = "",
    limit: int = Query(default=5000, ge=1, le=10000),
) -> Response:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    payload = players(
        date_from=date_from,
        date_to=date_to,
        location=location,
        organizer=organizer,
        level=level,
        format_value=format_value,
        tournament_type_value=tournament_type_value,
        time_period_value=time_period_value,
        rating_min=rating_min,
        rating_max=rating_max,
        search=search,
        limit=limit,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Players"
    ws.append(["player_name"])
    for item in payload["items"]:
        ws.append([item["player_name"]])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EEF7")
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 34

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return Response(
        stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="lunda_players_export.xlsx"'},
    )


@app.get("/api/tournaments")
def tournaments(
    view: str = Query(default="week", pattern="^(day|week)$"),
    date_value: str = Query(default="", alias="date"),
    date_from: str = "",
    date_to: str = "",
    location: list[str] | None = Query(default=None),
    organizer: list[str] | None = Query(default=None),
    level: list[str] | None = Query(default=None),
    format_value: list[str] | None = Query(default=None, alias="format"),
    tournament_type_value: list[str] | None = Query(default=None, alias="tournament_type"),
    time_period_value: list[str] | None = Query(default=None, alias="time_period"),
) -> dict[str, Any]:
    if view == "day":
        target = date_value or today_iso()
        date_from = date_to = target
    elif not date_from and not date_to:
        date_from, date_to = week_bounds(date_value or None)

    selected_levels = set(split_multi(level))
    selected_types = set(split_multi(tournament_type_value))
    selected_periods = set(split_multi(time_period_value))
    where_sql, params = tournament_filters(
        date_from=date_from,
        date_to=date_to,
        locations=split_multi(location),
        organizers=split_multi(organizer),
        formats=split_multi(format_value),
    )
    conn = connect()
    try:
        rows = conn.execute(
            f"""
            SELECT
                t.*,
                (SELECT COUNT(*) FROM final_participations fp WHERE fp.tournament_id = t.id) AS final_participant_count,
                (SELECT COUNT(*) FROM current_participants cp WHERE cp.tournament_id = t.id AND cp.active = 1) AS current_participant_count
            FROM tournaments t
            {where_sql}
            ORDER BY t.starts_at, t.location, t.organizer
            """,
            params,
        ).fetchall()
        items = [
            row_to_tournament(row)
            for row in rows
            if tournament_matches_extra_filters(
                row,
                selected_levels=selected_levels,
                selected_types=selected_types,
                selected_periods=selected_periods,
            )
        ]
        return {"items": items, "date_from": date_from, "date_to": date_to, "view": view}
    finally:
        conn.close()


HTML = r"""
<!doctype html>
<html lang="ru">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>Lunda Stats</title>
  <style>
    :root {
      color-scheme: light;
      --bg: #f6f7f9;
      --panel: #ffffff;
      --line: #d8dee8;
      --text: #1d2430;
      --muted: #667085;
      --accent: #0f766e;
    }
    * { box-sizing: border-box; }
    body { margin: 0; background: var(--bg); color: var(--text); font: 14px/1.4 Arial, sans-serif; }
    button, input, select { font: inherit; }
    .app { min-height: 100vh; display: grid; grid-template-rows: auto 1fr; }
    .topbar { background: #12202f; color: #fff; padding: 14px 24px; display: flex; align-items: center; justify-content: space-between; gap: 16px; }
    .brand { font-size: 20px; font-weight: 700; }
    .tabs { display: flex; gap: 8px; }
    .tab { border: 1px solid rgba(255,255,255,.25); background: transparent; color: #fff; padding: 8px 12px; border-radius: 6px; cursor: pointer; }
    .tab.active { background: #fff; color: #12202f; }
    .main { padding: 18px 24px 28px; display: grid; gap: 14px; }
    .filters { background: var(--panel); border-bottom: 1px solid var(--line); padding: 12px 24px; display: grid; grid-template-columns: repeat(8, minmax(120px, 1fr)); gap: 10px; align-items: end; }
    label { display: grid; gap: 5px; color: var(--muted); font-size: 12px; }
    input, select { min-height: 36px; border: 1px solid var(--line); border-radius: 6px; background: #fff; color: var(--text); padding: 7px 9px; }
    select[multiple] { height: 80px; }
    .actions { display: flex; gap: 8px; }
    .btn { border: 1px solid var(--line); background: #fff; padding: 8px 12px; border-radius: 6px; cursor: pointer; }
    .btn.primary { background: var(--accent); border-color: var(--accent); color: #fff; }
    .summary { display: flex; gap: 12px; color: var(--muted); align-items: center; }
    .panel { background: var(--panel); border: 1px solid var(--line); border-radius: 8px; overflow: hidden; }
    table { width: 100%; border-collapse: collapse; }
    th, td { border-bottom: 1px solid var(--line); padding: 9px 10px; text-align: left; vertical-align: top; }
    th { background: #eef2f7; color: #344054; font-size: 12px; position: sticky; top: 0; z-index: 2; }
    tr:hover td { background: #f8fafc; }
    .player-name { font-weight: 700; }
    .tags { display: flex; gap: 5px; flex-wrap: wrap; }
    .tag { display: inline-flex; padding: 2px 6px; border-radius: 5px; background: #eef2f7; color: #344054; font-size: 12px; }
    .details { color: var(--muted); font-size: 12px; margin-top: 3px; }
    .calendar-tools { display: flex; justify-content: space-between; align-items: center; gap: 12px; }
    .week-nav { display: flex; gap: 8px; align-items: center; }
    .calendar { display: grid; grid-template-columns: 58px repeat(7, 1fr); border-top: 1px solid var(--line); border-left: 1px solid var(--line); background: #fff; min-height: 720px; }
    .time-col, .day-col { position: relative; border-right: 1px solid var(--line); }
    .day-head, .time-head { height: 38px; border-bottom: 1px solid var(--line); background: #eef2f7; padding: 9px; font-weight: 700; }
    .hour-line { position: absolute; left: 0; right: 0; border-top: 1px solid #edf1f6; color: #98a2b3; font-size: 11px; padding-left: 6px; }
    .event { position: absolute; min-height: 28px; border-radius: 6px; color: #fff; padding: 5px 7px; overflow: hidden; cursor: pointer; box-shadow: 0 1px 4px rgba(0,0,0,.18); border: 1px solid rgba(255,255,255,.22); }
    .event-title { font-weight: 700; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; font-size: 12px; }
    .event-meta { opacity: .92; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; font-size: 11px; }
    .day-list { display: grid; grid-template-columns: repeat(auto-fill, minmax(280px, 1fr)); gap: 10px; }
    .event-row { border-left: 5px solid var(--accent); padding: 10px; background: #fff; border-radius: 8px; border-top: 1px solid var(--line); border-right: 1px solid var(--line); border-bottom: 1px solid var(--line); cursor: pointer; }
    .hidden { display: none; }
    dialog { width: min(720px, calc(100vw - 32px)); border: 1px solid var(--line); border-radius: 8px; padding: 0; }
    dialog::backdrop { background: rgba(16,24,40,.35); }
    .modal-head { padding: 14px 16px; border-bottom: 1px solid var(--line); display: flex; justify-content: space-between; gap: 12px; }
    .modal-body { padding: 16px; display: grid; gap: 8px; }
    .x { border: 0; background: transparent; font-size: 22px; cursor: pointer; }
    @media (max-width: 980px) {
      .filters { grid-template-columns: repeat(2, minmax(130px, 1fr)); }
      .calendar { grid-template-columns: 48px repeat(7, minmax(120px, 1fr)); overflow-x: auto; }
      .main { padding: 14px; }
      .topbar { padding: 12px 14px; align-items: start; flex-direction: column; }
    }
  </style>
</head>
<body>
  <div class="app">
    <header class="topbar">
      <div class="brand">Lunda Stats</div>
      <div class="tabs">
        <button class="tab active" data-tab="players">Игроки</button>
        <button class="tab" data-tab="schedule">Расписание</button>
      </div>
    </header>
    <section class="filters">
      <label>С даты<input id="dateFrom" type="date"></label>
      <label>По дату<input id="dateTo" type="date"></label>
      <label>Клуб<select id="locations" multiple></select></label>
      <label>Организатор<select id="organizers" multiple></select></label>
      <label>Уровень<select id="levels" multiple></select></label>
      <label>Формат<select id="formats" multiple></select></label>
      <label>Тип<select id="tournamentTypes" multiple></select></label>
      <label>Время<select id="timePeriods" multiple></select></label>
      <label id="searchLabel">Игрок<input id="search" placeholder="поиск по имени"></label>
      <label id="ratingMinLabel">Рейтинг от<input id="ratingMin" type="number" min="1" max="7" step="0.01" placeholder="1.00"></label>
      <label id="ratingMaxLabel">Рейтинг до<input id="ratingMax" type="number" min="1" max="7" step="0.01" placeholder="7.00"></label>
      <label id="viewLabel" class="hidden">Вид<select id="scheduleView"><option value="week">Неделя</option><option value="day">День</option></select></label>
      <div class="actions">
        <button class="btn primary" id="apply">Показать</button>
        <button class="btn" id="exportPlayers" type="button">Excel</button>
        <button class="btn" id="reset">Сброс</button>
      </div>
    </section>
    <main class="main">
      <section id="playersView">
        <div class="summary" id="playersSummary"></div>
        <div class="panel">
          <table>
            <thead><tr><th>Игрок</th><th>Рейтинг</th><th>Участий</th><th>Клубы</th><th>Организаторы</th><th>Уровни</th></tr></thead>
            <tbody id="playersBody"></tbody>
          </table>
        </div>
      </section>
      <section id="scheduleViewBox" class="hidden">
        <div class="calendar-tools">
          <div class="summary" id="scheduleSummary"></div>
          <div class="week-nav">
            <button class="btn" id="prevWeek" type="button">Назад</button>
            <button class="btn" id="thisWeek" type="button">Текущая</button>
            <button class="btn" id="nextWeek" type="button">Вперёд</button>
          </div>
        </div>
        <div id="weekCalendar" class="panel calendar"></div>
        <div id="dayList" class="day-list hidden"></div>
      </section>
    </main>
  </div>
  <dialog id="modal">
    <div class="modal-head"><strong id="modalTitle"></strong><button class="x" id="closeModal">×</button></div>
    <div class="modal-body" id="modalBody"></div>
  </dialog>
  <script>
    const state = { tab: "players", filters: null, tournaments: [] };
    const qs = (id) => document.getElementById(id);
    const selected = (id) => Array.from(qs(id).selectedOptions).map((option) => option.value);
    const pad = (n) => String(n).padStart(2, "0");
    const multiSelectIds = ["locations", "organizers", "levels", "formats", "tournamentTypes", "timePeriods"];

    function monday(value) {
      const d = value ? new Date(value + "T00:00:00") : new Date();
      const day = (d.getDay() + 6) % 7;
      d.setDate(d.getDate() - day);
      return d;
    }
    function isoDate(d) { return `${d.getFullYear()}-${pad(d.getMonth() + 1)}-${pad(d.getDate())}`; }
    function addDays(d, count) { const x = new Date(d); x.setDate(x.getDate() + count); return x; }
    function shiftSchedule(days) {
      const base = qs("dateFrom").value || isoDate(new Date());
      const shifted = addDays(monday(base), days);
      qs("dateFrom").value = isoDate(shifted);
      qs("dateTo").value = isoDate(addDays(shifted, 6));
      if (state.tab === "schedule") refresh();
    }
    function paramsBase() {
      const params = new URLSearchParams();
      if (qs("dateFrom").value) params.set("date_from", qs("dateFrom").value);
      if (qs("dateTo").value) params.set("date_to", qs("dateTo").value);
      for (const value of selected("locations")) params.append("location", value);
      for (const value of selected("organizers")) params.append("organizer", value);
      for (const value of selected("levels")) params.append("level", value);
      for (const value of selected("formats")) params.append("format", value);
      for (const value of selected("tournamentTypes")) params.append("tournament_type", value);
      for (const value of selected("timePeriods")) params.append("time_period", value);
      return params;
    }
    function fillSelect(id, values) {
      qs(id).innerHTML = values.map((value) => {
        const optionValue = typeof value === "object" ? value.value : value;
        const label = typeof value === "object" ? value.label : value;
        return `<option value="${escapeHtml(optionValue)}">${escapeHtml(label)}</option>`;
      }).join("");
    }
    function escapeHtml(value) {
      return String(value ?? "").replace(/[&<>"']/g, (char) => ({ "&": "&amp;", "<": "&lt;", ">": "&gt;", '"': "&quot;", "'": "&#39;" }[char]));
    }
    function tags(values) { return `<div class="tags">${values.map((v) => `<span class="tag">${escapeHtml(v)}</span>`).join("")}</div>`; }
    function apiUrl(path, params) {
      const url = new URL(path, window.location.href);
      url.username = "";
      url.password = "";
      if (params) url.search = params.toString();
      return url.toString();
    }
    async function loadFilters() {
      const res = await fetch(apiUrl("/api/filters"));
      state.filters = await res.json();
      fillSelect("locations", state.filters.locations);
      fillSelect("organizers", state.filters.organizers);
      fillSelect("levels", state.filters.levels);
      fillSelect("formats", state.filters.formats);
      fillSelect("tournamentTypes", state.filters.tournament_types);
      fillSelect("timePeriods", state.filters.time_periods);
      qs("dateFrom").value = state.filters.date_min || "";
      qs("dateTo").value = state.filters.date_max || "";
    }
    async function loadPlayers() {
      const params = paramsBase();
      if (qs("search").value.trim()) params.set("search", qs("search").value.trim());
      if (qs("ratingMin").value) params.set("rating_min", qs("ratingMin").value);
      if (qs("ratingMax").value) params.set("rating_max", qs("ratingMax").value);
      const res = await fetch(apiUrl("/api/players", params));
      const data = await res.json();
      qs("playersSummary").textContent = `Игроков: ${data.total}`;
      qs("playersBody").innerHTML = data.items.map((item) => `
        <tr>
          <td><div class="player-name">${escapeHtml(item.player_name)}</div><div class="details">${item.tournaments.slice(0, 3).map(t => `${t.date} ${t.time || ""} · ${t.title}`).join("<br>")}</div></td>
          <td>${item.latest_rating == null ? "" : Number(item.latest_rating).toFixed(2)}</td>
          <td>${item.tournament_count}</td>
          <td>${tags(item.locations)}</td>
          <td>${tags(item.organizers)}</td>
          <td>${tags(item.levels)}</td>
        </tr>
      `).join("");
    }
    function playerParams() {
      const params = paramsBase();
      if (qs("search").value.trim()) params.set("search", qs("search").value.trim());
      if (qs("ratingMin").value) params.set("rating_min", qs("ratingMin").value);
      if (qs("ratingMax").value) params.set("rating_max", qs("ratingMax").value);
      params.set("limit", "10000");
      return params;
    }
    async function loadSchedule() {
      const params = new URLSearchParams();
      const view = qs("scheduleView").value;
      params.set("view", view);
      if (qs("dateFrom").value) params.set("date", qs("dateFrom").value);
      for (const value of selected("locations")) params.append("location", value);
      for (const value of selected("organizers")) params.append("organizer", value);
      for (const value of selected("levels")) params.append("level", value);
      for (const value of selected("formats")) params.append("format", value);
      for (const value of selected("tournamentTypes")) params.append("tournament_type", value);
      for (const value of selected("timePeriods")) params.append("time_period", value);
      const res = await fetch(apiUrl("/api/tournaments", params));
      const data = await res.json();
      state.tournaments = data.items;
      qs("scheduleSummary").textContent = `Турниров: ${data.items.length} · ${data.date_from || ""} - ${data.date_to || ""}`;
      if (view === "day") renderDay(data.items); else renderWeek(data.items, data.date_from);
    }
    function minutes(value) {
      if (!value) return 0;
      const d = new Date(value);
      return d.getHours() * 60 + d.getMinutes();
    }
    function assignLanes(items) {
      const byDay = new Map();
      for (const item of items) {
        const day = item.tournament_date;
        if (!byDay.has(day)) byDay.set(day, []);
        byDay.get(day).push(item);
      }
      const laneById = new Map();
      for (const dayItems of byDay.values()) {
        const lanes = [];
        for (const item of dayItems.sort((a, b) => minutes(a.starts_at) - minutes(b.starts_at))) {
          const start = minutes(item.starts_at);
          const end = minutes(item.ends_at) || start + 120;
          let lane = lanes.findIndex((lastEnd) => lastEnd <= start);
          if (lane < 0) { lane = lanes.length; lanes.push(end); } else { lanes[lane] = end; }
          laneById.set(item.id, { lane, lanes: lanes.length });
        }
      }
      return laneById;
    }
    function renderWeek(items, startDate) {
      qs("weekCalendar").classList.remove("hidden");
      qs("dayList").classList.add("hidden");
      const start = monday(startDate);
      const days = Array.from({ length: 7 }, (_, i) => isoDate(addDays(start, i)));
      const laneById = assignLanes(items);
      const startMin = 7 * 60;
      const endMin = 24 * 60;
      const height = 900;
      let html = `<div class="time-col"><div class="time-head"></div>`;
      for (let h = 7; h <= 23; h++) html += `<div class="hour-line" style="top:${38 + ((h * 60 - startMin) / (endMin - startMin)) * height}px">${pad(h)}:00</div>`;
      html += `</div>`;
      for (const day of days) {
        html += `<div class="day-col"><div class="day-head">${day}</div>`;
        for (let h = 7; h <= 23; h++) html += `<div class="hour-line" style="top:${38 + ((h * 60 - startMin) / (endMin - startMin)) * height}px"></div>`;
        for (const item of items.filter((x) => x.tournament_date === day)) {
          const start = Math.max(startMin, minutes(item.starts_at));
          const end = Math.min(endMin, minutes(item.ends_at) || start + 120);
          const top = 38 + ((start - startMin) / (endMin - startMin)) * height;
          const blockHeight = Math.max(32, ((end - start) / (endMin - startMin)) * height);
          const lane = laneById.get(item.id) || { lane: 0, lanes: 1 };
          const width = 94 / lane.lanes;
          const left = 3 + lane.lane * width;
          html += `<div class="event" data-id="${item.id}" style="background:${item.color}; top:${top}px; height:${blockHeight}px; left:${left}%; width:${width - 2}%">
            <div class="event-title">${escapeHtml(item.organizer || "")}</div>
            <div class="event-meta">${escapeHtml(item.starts_at_display)} ${escapeHtml(item.format || "")}</div>
            <div class="event-meta">${escapeHtml(item.tournament_type_label || "")} · ${escapeHtml(item.skill_level || "")} · ${escapeHtml(item.location || "")}</div>
          </div>`;
        }
        html += `</div>`;
      }
      qs("weekCalendar").innerHTML = html;
      qs("weekCalendar").style.minHeight = `${height + 39}px`;
      bindEvents();
    }
    function renderDay(items) {
      qs("weekCalendar").classList.add("hidden");
      qs("dayList").classList.remove("hidden");
      qs("dayList").innerHTML = items.map((item) => `
        <div class="event-row" data-id="${item.id}" style="border-left-color:${item.color}">
          <strong>${escapeHtml(item.starts_at_display)} - ${escapeHtml(item.ends_at_display)} · ${escapeHtml(item.title)}</strong>
          <div class="details">${escapeHtml(item.organizer)} · ${escapeHtml(item.location)} · ${escapeHtml(item.tournament_type_label)} · ${escapeHtml(item.time_period_label)}</div>
          <div class="details">${escapeHtml(item.skill_level)} · ${escapeHtml(item.format)}</div>
          <div class="details">${escapeHtml(item.price_label || "")} · ${item.final_participant_count || item.current_participant_count || 0} участников</div>
        </div>
      `).join("");
      bindEvents();
    }
    function bindEvents() {
      document.querySelectorAll("[data-id]").forEach((node) => {
        node.addEventListener("click", () => showTournament(Number(node.dataset.id)));
      });
    }
    function showTournament(id) {
      const item = state.tournaments.find((x) => x.id === id);
      if (!item) return;
      qs("modalTitle").textContent = item.title;
      qs("modalBody").innerHTML = `
        <div><strong>Время:</strong> ${escapeHtml(item.tournament_date)} ${escapeHtml(item.time_label || "")}</div>
        <div><strong>Организатор:</strong> ${escapeHtml(item.organizer || "")}</div>
        <div><strong>Место:</strong> ${escapeHtml(item.location || "")}</div>
        <div><strong>Тип:</strong> ${escapeHtml(item.tournament_type_label || "")}</div>
        <div><strong>Время дня:</strong> ${escapeHtml(item.time_period_label || "")}</div>
        <div><strong>Уровень:</strong> ${escapeHtml(item.skill_level || "")}</div>
        <div><strong>Формат:</strong> ${escapeHtml(item.format || "")}</div>
        <div><strong>Стоимость:</strong> ${escapeHtml(item.price_label || "")}</div>
        <div><strong>Заявка:</strong> ${escapeHtml((item.participants_current || "") + "/" + (item.participants_capacity || "") + " " + (item.participants_unit || ""))}</div>
        <div><strong>Финальных участников:</strong> ${item.final_participant_count || 0}</div>
        <div><strong>Статус:</strong> ${escapeHtml(item.source_status || "")}</div>
      `;
      qs("modal").showModal();
    }
    function setTab(tab) {
      state.tab = tab;
      document.querySelectorAll(".tab").forEach((btn) => btn.classList.toggle("active", btn.dataset.tab === tab));
      qs("playersView").classList.toggle("hidden", tab !== "players");
      qs("scheduleViewBox").classList.toggle("hidden", tab !== "schedule");
      qs("searchLabel").classList.toggle("hidden", tab !== "players");
      qs("ratingMinLabel").classList.toggle("hidden", tab !== "players");
      qs("ratingMaxLabel").classList.toggle("hidden", tab !== "players");
      qs("exportPlayers").classList.toggle("hidden", tab !== "players");
      qs("viewLabel").classList.toggle("hidden", tab !== "schedule");
      document.querySelector(".week-nav").classList.toggle("hidden", tab !== "schedule");
      refresh();
    }
    async function refresh() {
      if (state.tab === "players") await loadPlayers();
      else await loadSchedule();
    }
    qs("apply").addEventListener("click", refresh);
    qs("exportPlayers").addEventListener("click", () => {
      window.location.href = apiUrl("/api/players/export", playerParams());
    });
    qs("prevWeek").addEventListener("click", () => shiftSchedule(-7));
    qs("thisWeek").addEventListener("click", () => {
      const start = monday();
      qs("dateFrom").value = isoDate(start);
      qs("dateTo").value = isoDate(addDays(start, 6));
      if (state.tab === "schedule") refresh();
    });
    qs("nextWeek").addEventListener("click", () => shiftSchedule(7));
    qs("reset").addEventListener("click", () => {
      document.querySelectorAll("select").forEach((select) => Array.from(select.options).forEach((option) => option.selected = false));
      qs("search").value = "";
      qs("ratingMin").value = "";
      qs("ratingMax").value = "";
      qs("dateFrom").value = state.filters?.date_min || "";
      qs("dateTo").value = state.filters?.date_max || "";
      qs("scheduleView").value = "week";
      refresh();
    });
    qs("closeModal").addEventListener("click", () => qs("modal").close());
    document.querySelectorAll(".tab").forEach((btn) => btn.addEventListener("click", () => setTab(btn.dataset.tab)));
    for (const id of multiSelectIds) {
      const select = qs(id);
      let beforeMouseDown = new Set();
      select.addEventListener("mousedown", () => {
        beforeMouseDown = new Set(selected(id));
      });
      select.addEventListener("click", (event) => {
        if (event.target.tagName !== "OPTION") return;
        const option = event.target;
        if (!beforeMouseDown.has(option.value)) return;
        event.preventDefault();
        option.selected = false;
        select.dispatchEvent(new Event("change", { bubbles: true }));
      });
    }
    loadFilters().then(refresh);
  </script>
</body>
</html>
"""
