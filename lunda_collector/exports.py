from __future__ import annotations

import sqlite3
from pathlib import Path
from typing import Any

from storage import player_rankings, schedule_rows, tournament_summary


def export_workbook(
    conn: sqlite3.Connection,
    output_path: str | Path,
    *,
    date_from: str = "",
    date_to: str = "",
    organizer: str = "",
    location: str = "",
    skill_level: str = "",
    format_value: str = "",
) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    output = Path(output_path).expanduser()
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Summary"
    _append_table(ws_summary, ["metric", "value"], sorted(tournament_summary(conn).items()))

    rankings = player_rankings(
        conn,
        date_from=date_from,
        date_to=date_to,
        organizer=organizer,
        location=location,
        skill_level=skill_level,
        format_value=format_value,
        limit=10000,
    )
    ws_players = wb.create_sheet("Players")
    _append_table(
        ws_players,
        ["player_name", "latest_rating", "tournament_count", "locations", "organizers", "levels"],
        [
            [
                row["player_name"],
                row["latest_rating"] if row["latest_rating"] is not None else "",
                row["tournament_count"],
                row["locations"] or "",
                row["organizers"] or "",
                row["levels"] or "",
            ]
            for row in rankings
        ],
    )

    schedule = schedule_rows(
        conn,
        date_from=date_from,
        date_to=date_to,
        organizer=organizer,
        location=location,
        skill_level=skill_level,
        format_value=format_value,
    )
    ws_schedule = wb.create_sheet("Schedule")
    _append_table(
        ws_schedule,
        [
            "id",
            "date",
            "time",
            "title",
            "organizer",
            "location",
            "level",
            "format",
            "price",
            "participants",
            "status",
        ],
        [
            [
                row["id"],
                row["tournament_date"] or row["date_label"] or "",
                row["time_label"] or "",
                row["title"] or "",
                row["organizer"] or "",
                row["location"] or "",
                row["skill_level"] or "",
                row["format"] or "",
                row["price_label"] or "",
                _participants_label(row),
                row["source_status"],
            ]
            for row in schedule
        ],
    )

    ws_final = wb.create_sheet("Final participations")
    final_rows = conn.execute(
        """
        SELECT
            t.tournament_date,
            t.time_label,
            t.title,
            t.organizer,
            t.location,
            COALESCE(p.display_name, fp.raw_name) AS player_name,
            fp.rating,
            fp.resolve_status,
            fp.finalized_at
        FROM final_participations fp
        JOIN tournaments t ON t.id = fp.tournament_id
        LEFT JOIN players p ON p.id = fp.player_id
        ORDER BY t.starts_at, t.organizer, t.title, player_name
        """
    ).fetchall()
    _append_table(
        ws_final,
        ["date", "time", "title", "organizer", "location", "player", "rating", "resolve_status", "finalized_at"],
        [[row[column] or "" for column in row.keys()] for row in final_rows],
    )

    ws_pending = wb.create_sheet("Pending")
    pending_rows = conn.execute(
        """
        SELECT
            pp.id,
            pp.raw_name,
            pp.normalized_name,
            t.title,
            t.organizer,
            t.date_label,
            t.time_label,
            pp.candidates_json,
            pp.updated_at
        FROM pending_players pp
        LEFT JOIN tournaments t ON t.id = pp.tournament_id
        WHERE pp.status = 'pending'
        ORDER BY pp.updated_at DESC
        """
    ).fetchall()
    _append_table(
        ws_pending,
        ["id", "raw_name", "normalized_name", "title", "organizer", "date", "time", "candidates_json", "updated_at"],
        [[row[column] or "" for column in row.keys()] for row in pending_rows],
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E8EEF7")
        _autosize(ws)

    wb.save(output)
    return output


def _append_table(ws: Any, headers: list[str], rows: list[Any]) -> None:
    ws.append(headers)
    for row in rows:
        if isinstance(row, tuple):
            ws.append(list(row))
        elif isinstance(row, list):
            ws.append(row)
        else:
            ws.append([row])


def _autosize(ws: Any) -> None:
    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 10), 80)


def _participants_label(row: sqlite3.Row) -> str:
    current = row["participants_current"]
    capacity = row["participants_capacity"]
    unit = row["participants_unit"] or ""
    if current is None or capacity is None:
        return ""
    return f"{current}/{capacity} {unit}".strip()
