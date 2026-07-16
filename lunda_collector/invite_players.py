#!/usr/bin/env python3
from __future__ import annotations

import argparse
import base64
import json
import os
import re
import subprocess
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from date_parser import MSK, iso_or_empty, parse_tournament_datetime
from live_phone_cycle import (
    LiveContext,
    create_backend,
    get_screen_size,
    launch_lunda,
    load_env_file,
    select_adb_device,
    tap_lunda_bottom_nav,
)
from storage import DEFAULT_DB_PATH, connect, init_db
from visible_cards import OCRLine, extract_ocr_lines


ADB_KEYBOARD_IME = "com.android.adbkeyboard/.AdbIME"


@dataclass(frozen=True)
class MyTournament:
    identity_key: str
    date_label: str
    time_label: str
    starts_at: str
    ends_at: str
    location: str
    participants_label: str
    tap_x: int
    tap_y: int
    raw: dict[str, Any]


def main() -> int:
    parser = argparse.ArgumentParser(description="Lunda player invitation helper")
    parser.add_argument("--db", default=os.environ.get("LUNDA_DB_PATH", str(DEFAULT_DB_PATH)))
    parser.add_argument("--parser-dir", default=os.environ.get("LUNDA_OLD_PARSER_DIR", ""))
    parser.add_argument("--device", default=os.environ.get("PHONE_ADB_HOST", ""))
    parser.add_argument("--out-dir", default="")
    parser.add_argument("--env-file", default="/opt/lunda-collector/app/.env")
    subparsers = parser.add_subparsers(dest="command", required=True)

    collect = subparsers.add_parser("collect-my-tournaments", help="Collect tournaments from Calendar of my events")
    collect.add_argument("--launch", action="store_true")
    collect.add_argument("--screens", type=int, default=20)
    collect.add_argument("--scroll-pixels", type=int, default=620)

    create_job = subparsers.add_parser("create-job", help="Create invite job from Excel file")
    create_job.add_argument("--my-tournament-id", type=int, required=True)
    create_job.add_argument("--xlsx", required=True)
    create_job.add_argument("--scheduled-for", default="")

    run_job = subparsers.add_parser("run-job", help="Run one queued invite job")
    run_job.add_argument("--job-id", type=int, required=True)
    run_job.add_argument("--launch", action="store_true")

    run_pending = subparsers.add_parser("run-pending", help="Run queued invite jobs")
    run_pending.add_argument("--launch", action="store_true")
    run_pending.add_argument("--limit", type=int, default=1)

    args = parser.parse_args()
    load_env_file(Path(args.env_file))

    conn = connect(args.db)
    init_db(conn)
    init_invite_db(conn)

    if args.command == "create-job":
        names = read_player_names_from_xlsx(Path(args.xlsx))
        job_id = create_invite_job(
            conn,
            my_tournament_id=args.my_tournament_id,
            player_names=names,
            scheduled_for=args.scheduled_for,
        )
        print(f"Invite job #{job_id} queued; players={len(names)}")
        return 0

    if args.command in {"collect-my-tournaments", "run-job", "run-pending"}:
        selected_device = args.device or select_adb_device()
        if not selected_device:
            print("ADB devices: none", file=sys.stderr)
            return 2
        parser_dir = Path(args.parser_dir).expanduser().resolve() if args.parser_dir else Path("")
        android, ocr = create_backend(parser_dir, selected_device)
        out_dir = Path(args.out_dir).expanduser().resolve() if args.out_dir else Path.cwd() / "work" / f"invite_{args.command}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        out_dir.mkdir(parents=True, exist_ok=True)
        ctx = LiveContext(android, ocr, out_dir)
        if args.launch:
            launch_lunda(selected_device, adb_path=getattr(android, "adb_path", "adb"))
            time.sleep(3.0)

    if args.command == "collect-my-tournaments":
        tournaments = collect_my_tournaments(conn, ctx, screens=args.screens, scroll_pixels=args.scroll_pixels)
        print(f"Collected my tournaments: {len(tournaments)}")
        print(f"Artifacts: {ctx.out_dir}")
        return 0

    if args.command == "run-job":
        return run_invite_job(conn, ctx, job_id=args.job_id)

    if args.command == "run-pending":
        count = 0
        for row in queued_invite_jobs(conn, limit=args.limit):
            result = run_invite_job(conn, ctx, job_id=int(row["id"]))
            count += 1
            if result != 0:
                return result
        print(f"Invite jobs processed: {count}")
        return 0

    return 2


def init_invite_db(conn) -> None:
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS my_tournaments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            identity_key TEXT NOT NULL UNIQUE,
            date_label TEXT,
            time_label TEXT,
            starts_at TEXT,
            ends_at TEXT,
            location TEXT,
            participants_label TEXT,
            tap_x INTEGER,
            tap_y INTEGER,
            raw_json TEXT NOT NULL DEFAULT '{}',
            source_status TEXT NOT NULL DEFAULT 'active',
            first_seen_at TEXT NOT NULL,
            last_seen_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS invite_jobs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            my_tournament_id INTEGER NOT NULL REFERENCES my_tournaments(id) ON DELETE CASCADE,
            status TEXT NOT NULL DEFAULT 'queued',
            created_at TEXT NOT NULL,
            scheduled_for TEXT,
            started_at TEXT,
            finished_at TEXT,
            player_count INTEGER NOT NULL DEFAULT 0,
            players_json TEXT NOT NULL DEFAULT '[]',
            stats_json TEXT NOT NULL DEFAULT '{}',
            error TEXT,
            log_path TEXT
        );
        """
    )
    conn.commit()


def read_player_names_from_xlsx(path: Path) -> list[str]:
    from openpyxl import load_workbook

    workbook = load_workbook(path.expanduser(), read_only=True, data_only=True)
    sheet = workbook.active
    names: list[str] = []
    for row in sheet.iter_rows(values_only=True):
        if not row:
            continue
        value = str(row[0] or "").strip()
        if not value:
            continue
        if value.lower() in {"player_name", "name", "имя", "фио", "игрок", "player"}:
            continue
        normalized = re.sub(r"\s+", " ", value)
        if normalized not in names:
            names.append(normalized)
    return names


def create_invite_job(conn, *, my_tournament_id: int, player_names: list[str], scheduled_for: str = "") -> int:
    now = now_iso()
    conn.execute(
        """
        INSERT INTO invite_jobs (
            my_tournament_id, status, created_at, scheduled_for,
            player_count, players_json
        )
        VALUES (?, 'queued', ?, ?, ?, ?)
        """,
        (my_tournament_id, now, scheduled_for, len(player_names), json.dumps(player_names, ensure_ascii=False)),
    )
    conn.commit()
    return int(conn.execute("SELECT last_insert_rowid()").fetchone()[0])


def queued_invite_jobs(conn, *, limit: int = 1):
    return conn.execute(
        """
        SELECT *
        FROM invite_jobs
        WHERE status = 'queued'
          AND (scheduled_for IS NULL OR scheduled_for = '' OR scheduled_for <= ?)
        ORDER BY created_at
        LIMIT ?
        """,
        (now_iso(), limit),
    ).fetchall()


def collect_my_tournaments(conn, ctx: LiveContext, *, screens: int, scroll_pixels: int) -> list[MyTournament]:
    if not open_my_events_screen(ctx):
        raise RuntimeError("my events screen was not reached")

    observations: list[MyTournament] = []
    previous_signature = ""
    repeated = 0
    for idx in range(screens):
        captured = ctx.capture(f"my_events_{idx + 1:02d}")
        if not captured:
            continue
        ocr_result, text, _ = captured
        if "календарь моих событий" not in text.lower():
            raise RuntimeError("unexpected screen while collecting my events")
        cards = parse_my_tournament_cards(ocr_result)
        observations.extend(cards)
        signature = "|".join(card.identity_key for card in cards)
        repeated = repeated + 1 if signature and signature == previous_signature else 0
        previous_signature = signature or previous_signature
        if repeated >= 2:
            break
        ctx.android.scroll_down(pixels=scroll_pixels)
        time.sleep(1.0)

    merged = merge_my_tournaments(observations)
    upsert_my_tournaments(conn, merged)
    (ctx.out_dir / "my_tournaments.json").write_text(
        json.dumps([card.__dict__ for card in merged], ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return merged


def open_my_events_screen(ctx: LiveContext, *, force_home: bool = False) -> bool:
    if force_home:
        if not ensure_main_home_screen(ctx):
            return False
    elif not ensure_home_screen(ctx):
        return False
    for attempt in range(3):
        captured = ctx.capture(f"open_my_events_{attempt + 1}")
        if not captured:
            continue
        ocr_result, text, _ = captured
        if is_my_events_screen(text):
            return True
        if is_main_home_screen(text):
            width, height = get_screen_size(ctx)
            ctx.android.tap(int(width * 0.93), int(height * 0.225))
            time.sleep(2.0)
            continue
        coords = find_line_center(ocr_result, "календарь моих событий") or find_line_center(ocr_result, "календарь моих")
        if coords:
            ctx.android.tap(coords["x"], coords["y"])
        else:
            width, height = get_screen_size(ctx)
            ctx.android.tap(width // 2, int(height * 0.30))
        time.sleep(2.0)
    captured = ctx.capture("my_events_check")
    return bool(captured and is_my_events_screen(captured[1]))


def ensure_main_home_screen(ctx: LiveContext) -> bool:
    for _ in range(8):
        captured = ctx.capture("ensure_main_home_for_invites")
        if not captured:
            continue
        _, text, _ = captured
        lower = text.lower()
        if is_main_home_screen(text):
            return True
        if "главная" in lower and "играть" in lower:
            tap_lunda_bottom_nav(ctx, "играть")
            time.sleep(1.5)
            tap_lunda_bottom_nav(ctx, "главная")
            time.sleep(2.0)
            continue
        ctx.android.go_back()
        time.sleep(1.0)
    return False


def ensure_home_screen(ctx: LiveContext) -> bool:
    for _ in range(8):
        captured = ctx.capture("ensure_home_for_invites")
        if not captured:
            continue
        _, text, _ = captured
        lower = text.lower()
        if is_my_events_screen(text):
            return True
        if is_main_home_screen(text):
            return True
        if "главная" in lower and "играть" in lower:
            tap_lunda_bottom_nav(ctx, "главная")
            time.sleep(2.0)
            continue
        ctx.android.go_back()
        time.sleep(1.0)
    return False


def has_my_events_header(text: str) -> bool:
    lower = text.lower()
    return "календарь моих событий" in lower or ("календарь моих" in lower and "событий" in lower)


def is_main_home_screen(text: str) -> bool:
    lower = text.lower()
    return "ваш город" in lower and has_my_events_header(lower)


def is_my_events_screen(text: str) -> bool:
    lower = text.lower()
    if "главная" in lower and "играть" in lower and "рейтинг" in lower:
        return False
    return has_my_events_header(lower) and "бронирования" in lower and ("июль" in lower or "турнир" in lower)


def parse_my_tournament_cards(ocr_result: dict[str, Any]) -> list[MyTournament]:
    lines = extract_ocr_lines(ocr_result)
    cards: list[MyTournament] = []
    for idx, line in enumerate(lines):
        date_label, time_label = split_event_datetime(line.text)
        if not date_label or not time_label:
            continue
        location_line = _next_line(lines, idx, lambda value: "|" in value or "padel" in value.lower() or "падел" in value.lower())
        participants_line = _next_line(lines, idx, lambda value: bool(re.search(r"\d+\s*/\s*\d+\s+(?:игрок|команд)", value.lower())))
        if not location_line or not participants_line:
            continue
        location = cleanup_my_event_location(join_wrapped_location(lines, location_line))
        parsed = parse_tournament_datetime(date_label, time_label)
        starts_at = iso_or_empty(parsed.starts_at)
        ends_at = iso_or_empty(parsed.ends_at)
        identity_key = "|".join(["my", starts_at, location.lower()])
        y_values = [line.y_min, line.y_max, location_line.y_min, location_line.y_max]
        if participants_line:
            y_values.extend([participants_line.y_min, participants_line.y_max])
        cards.append(
            MyTournament(
                identity_key=identity_key,
                date_label=date_label,
                time_label=time_label,
                starts_at=starts_at,
                ends_at=ends_at,
                location=location,
                participants_label=participants_line.text if participants_line else "",
                tap_x=(line.x_min + line.x_max) // 2,
                tap_y=(min(y_values) + max(y_values)) // 2,
                raw={
                    "datetime": line.text,
                    "location": location_line.text,
                    "participants": participants_line.text if participants_line else "",
                },
            )
        )
    return cards


def split_event_datetime(text: str) -> tuple[str, str]:
    match = re.search(
        r"((?:пн|вт|ср|чт|пт|сб|вс)\s+\d{1,2}\s+[а-яё]+)\s*\|?\s*((?:\d{1,2}:\d{2})\s*[-–—]\s*(?:\d{1,2}:\d{2}))",
        text,
        re.IGNORECASE,
    )
    if not match:
        return "", ""
    return match.group(1), match.group(2).replace("–", "-").replace("—", "-")


def cleanup_my_event_location(text: str) -> str:
    value = re.sub(r"^[^\wА-Яа-яЁё]+", "", text).strip()
    value = re.sub(r"\s+", " ", value)
    value = value.replace("Санкт- Петербург", "Санкт-Петербург")
    value = value.replace("санкт- петербург", "Санкт-Петербург")
    return value


def join_wrapped_location(lines: list[OCRLine], location_line: OCRLine) -> str:
    parts = [location_line.text]
    for line in lines:
        if line.y_min <= location_line.y_min:
            continue
        if line.y_min - location_line.y_max > 70:
            break
        lower = line.text.lower().strip()
        if lower.startswith("петербург") or lower.startswith("санкт-петербург"):
            parts.append(line.text)
            break
    return " ".join(parts)


def _next_line(lines: list[OCRLine], idx: int, predicate) -> OCRLine | None:
    base = lines[idx]
    for candidate in lines[idx + 1 : idx + 6]:
        if candidate.y_min - base.y_min > 240:
            break
        if predicate(candidate.text):
            return candidate
    return None


def merge_my_tournaments(cards: list[MyTournament]) -> list[MyTournament]:
    merged: dict[str, MyTournament] = {}
    for card in cards:
        if not card.starts_at or not card.location:
            continue
        merged[card.identity_key] = card
    return sorted(merged.values(), key=lambda item: (item.starts_at, item.location))


def upsert_my_tournaments(conn, cards: list[MyTournament]) -> None:
    now = now_iso()
    seen = set()
    for card in cards:
        seen.add(card.identity_key)
        conn.execute(
            """
            INSERT INTO my_tournaments (
                identity_key, date_label, time_label, starts_at, ends_at,
                location, participants_label, tap_x, tap_y, raw_json,
                source_status, first_seen_at, last_seen_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'active', ?, ?)
            ON CONFLICT(identity_key) DO UPDATE SET
                date_label = excluded.date_label,
                time_label = excluded.time_label,
                starts_at = excluded.starts_at,
                ends_at = excluded.ends_at,
                location = excluded.location,
                participants_label = excluded.participants_label,
                tap_x = excluded.tap_x,
                tap_y = excluded.tap_y,
                raw_json = excluded.raw_json,
                source_status = 'active',
                last_seen_at = excluded.last_seen_at
            """,
            (
                card.identity_key,
                card.date_label,
                card.time_label,
                card.starts_at,
                card.ends_at,
                card.location,
                card.participants_label,
                card.tap_x,
                card.tap_y,
                json.dumps(card.raw, ensure_ascii=False),
                now,
                now,
            ),
        )
    if seen:
        placeholders = ",".join("?" for _ in seen)
        conn.execute(
            f"UPDATE my_tournaments SET source_status = 'missing', last_seen_at = ? WHERE identity_key NOT IN ({placeholders})",
            [now, *sorted(seen)],
        )
    conn.commit()


def run_invite_job(conn, ctx: LiveContext, *, job_id: int) -> int:
    row = conn.execute(
        """
        SELECT ij.*, mt.identity_key, mt.starts_at, mt.location, mt.time_label
        FROM invite_jobs ij
        JOIN my_tournaments mt ON mt.id = ij.my_tournament_id
        WHERE ij.id = ?
        """,
        (job_id,),
    ).fetchone()
    if not row:
        print(f"Invite job #{job_id} not found", file=sys.stderr)
        return 2
    if row["status"] not in {"queued", "error"}:
        print(f"Invite job #{job_id} status={row['status']}; skipping")
        return 0

    log_path = ctx.out_dir / f"invite_job_{job_id}.log"
    players = json.loads(row["players_json"] or "[]")
    stats = {"players": len(players), "selected": 0, "not_found": 0, "errors": 0}
    conn.execute(
        "UPDATE invite_jobs SET status='running', started_at=?, log_path=?, error=NULL WHERE id=?",
        (now_iso(), str(log_path), job_id),
    )
    conn.commit()

    try:
        ensure_adb_keyboard(ctx)
        if not open_my_tournament_detail(ctx, row):
            raise RuntimeError("target my tournament was not opened")
        if not open_invite_screen(ctx):
            raise RuntimeError("invite screen was not opened")
        disable_only_partners(ctx)
        for index, name in enumerate(players):
            result = invite_one_player(ctx, str(name))
            stats[result] = stats.get(result, 0) + 1
            log_path.write_text(json.dumps(stats, ensure_ascii=False, indent=2), encoding="utf-8")
            if result != "selected" or index < len(players) - 1:
                clear_search_field(ctx)
        if stats["selected"] > 0 and not tap_final_invite_button(ctx):
            raise RuntimeError("final invite button was not tapped")
        conn.execute(
            "UPDATE invite_jobs SET status='done', finished_at=?, stats_json=? WHERE id=?",
            (now_iso(), json.dumps(stats, ensure_ascii=False), job_id),
        )
        conn.commit()
        print(json.dumps(stats, ensure_ascii=False, indent=2))
        return 0
    except Exception as exc:
        stats["error"] = str(exc)
        conn.execute(
            "UPDATE invite_jobs SET status='error', finished_at=?, stats_json=?, error=? WHERE id=?",
            (now_iso(), json.dumps(stats, ensure_ascii=False), str(exc), job_id),
        )
        conn.commit()
        print(f"Invite job #{job_id} failed: {exc}", file=sys.stderr)
        return 1


def open_my_tournament_detail(ctx: LiveContext, tournament_row) -> bool:
    if not open_my_events_screen(ctx, force_home=True):
        return False
    target_key = str(tournament_row["identity_key"])
    target_starts_at = str(tournament_row["starts_at"] or "")
    target_location = normalize_invite_location(str(tournament_row["location"] or ""))
    _, height = get_screen_size(ctx)
    min_card_y = int(height * 0.55)
    for _ in range(20):
        captured = ctx.capture("find_my_tournament")
        if not captured:
            continue
        ocr_result, _, _ = captured
        cards = parse_my_tournament_cards(ocr_result)
        for card in cards:
            if card.tap_y < min_card_y:
                continue
            if card.identity_key == target_key or (
                card.starts_at == target_starts_at
                and normalize_invite_location(card.location) == target_location
            ):
                ctx.android.tap(card.tap_x, card.tap_y)
                time.sleep(2.0)
                return True
        scroll_my_events_list_down(ctx, pixels=620)
        time.sleep(1.0)
    return False


def scroll_my_events_to_top(ctx: LiveContext) -> None:
    for _ in range(5):
        ctx.android.scroll_down(pixels=-850)
        time.sleep(0.6)


def scroll_my_events_list_down(ctx: LiveContext, *, pixels: int = 620) -> bool:
    width, height = get_screen_size(ctx)
    start_y = int(height * 0.88)
    end_y = max(int(height * 0.56), start_y - pixels)
    adb_path = getattr(ctx.android, "adb_path", "adb")
    device_id = getattr(ctx.android, "device_id", "") or select_adb_device()
    result = subprocess.run(
        [
            adb_path,
            "-s",
            device_id,
            "shell",
            "input",
            "swipe",
            str(width // 2),
            str(start_y),
            str(width // 2),
            str(end_y),
            "650",
        ],
        capture_output=True,
        text=True,
        timeout=10,
    )
    time.sleep(0.5)
    return result.returncode == 0


def normalize_invite_location(value: str) -> str:
    normalized = value.lower().replace("санкт- петербург", "санкт-петербург")
    normalized = re.sub(r"\s+", " ", normalized).strip()
    return normalized


def open_invite_screen(ctx: LiveContext) -> bool:
    for attempt in range(12):
        captured = ctx.capture(f"open_invite_button_{attempt + 1}")
        if not captured:
            continue
        ocr_result, text, _ = captured
        if is_invite_screen(text):
            return True
        coords = find_line_center(ocr_result, "пригласить игроков")
        if coords:
            ctx.android.tap(coords["x"], coords["y"])
            time.sleep(2.0)
            continue
        ctx.android.scroll_down(pixels=420)
        time.sleep(1.0)
    return False


def is_invite_screen(text: str) -> bool:
    lower = text.lower()
    return "пригласить игроков" in lower and "только мои напарники" in lower and "найти игрока" in lower


def disable_only_partners(ctx: LiveContext) -> None:
    captured = ctx.capture("invite_toggle_check")
    if not captured:
        return
    ocr_result, text, _ = captured
    if "игроки в других городах" in text.lower():
        return
    coords = find_line_center(ocr_result, "только мои напарники")
    width, height = get_screen_size(ctx)
    tap_y = coords["y"] if coords else int(height * 0.335)
    ctx.android.tap(int(width * 0.90), tap_y)
    time.sleep(1.5)


def invite_one_player(ctx: LiveContext, name: str) -> str:
    if not focus_search_field(ctx):
        return "errors"
    send_unicode_text(ctx, name)
    time.sleep(2.0)
    selected = select_search_results(ctx)
    return "selected" if selected else "not_found"


def focus_search_field(ctx: LiveContext) -> bool:
    captured = ctx.capture("invite_focus_search")
    if not captured:
        return False
    ocr_result, _, _ = captured
    coords = find_line_center(ocr_result, "найти игрока")
    width, height = get_screen_size(ctx)
    if coords:
        ctx.android.tap(max(coords["x"], int(width * 0.42)), coords["y"])
    else:
        ctx.android.tap(width // 2, int(height * 0.35))
    time.sleep(0.8)
    return True


def clear_search_field(ctx: LiveContext) -> None:
    captured = ctx.capture("invite_clear_search")
    width, height = get_screen_size(ctx)
    if captured:
        ocr_result, _, _ = captured
        y = search_field_y(ocr_result)
        if y:
            ctx.android.tap(int(width * 0.45), y)
            time.sleep(0.3)
            delete_search_text(ctx, repeats=45)
            time.sleep(0.7)
            return
    ctx.android.tap(width // 2, int(height * 0.42))
    time.sleep(0.3)
    delete_search_text(ctx, repeats=45)
    time.sleep(0.5)


def select_search_results(ctx: LiveContext) -> int:
    captured = ctx.capture("invite_search_results")
    if not captured:
        return 0
    ocr_result, text, _ = captured
    lower = text.lower()
    if "нет игроков" in lower or "изменить город" in lower:
        return 0
    lines = extract_ocr_lines(ocr_result)
    result_lines = candidate_player_result_lines(lines)
    selected = 0
    width, _ = get_screen_size(ctx)
    for y in candidate_player_result_tap_ys(result_lines):
        ctx.android.tap(width // 2, y)
        selected += 1
        time.sleep(0.4)
    return selected


def candidate_player_result_tap_ys(lines: list[OCRLine]) -> list[int]:
    tap_ys: list[int] = []
    for line in sorted(lines, key=lambda item: (item.y_min + item.y_max) // 2):
        y = (line.y_min + line.y_max) // 2
        if tap_ys and y - tap_ys[-1] < 90:
            continue
        tap_ys.append(y)
    return tap_ys


def candidate_player_result_lines(lines: list[OCRLine]) -> list[OCRLine]:
    results: list[OCRLine] = []
    active = False
    for line in lines:
        text = line.text.strip()
        lower = text.lower()
        if "игроки рядом" in lower or lower == "игроки":
            active = True
            continue
        if "игроки в других городах" in lower:
            break
        if not active:
            continue
        if (
            lower in {"рядом", "пригласить"}
            or "adb keyboard" in lower
            or re.fullmatch(r"\(?[A-ZА-ЯЁ]{1,3}\)?", text)
        ):
            continue
        if "санкт-петербург" in lower or re.search(r"\b[lr]+\s*\|\s*\d", lower):
            continue
        if re.search(r"[A-Za-zА-Яа-яЁё]{2,}", text) and line.x_min < 350:
            results.append(line)
    return results


def tap_final_invite_button(ctx: LiveContext) -> bool:
    width, height = get_screen_size(ctx)
    captured = ctx.capture("invite_final_button_before_tap")
    if captured:
        ocr_result, _, _ = captured
        coords = find_bottom_line_center(ocr_result, "пригласить")
        if coords:
            ctx.android.tap(coords["x"], coords["y"])
        else:
            ctx.android.tap(width // 2, int(height * 0.88))
    else:
        ctx.android.tap(width // 2, int(height * 0.88))
    time.sleep(2.0)
    ctx.capture("invite_final_button_after_tap")
    return True


def ensure_adb_keyboard(ctx: LiveContext) -> None:
    adb_path = getattr(ctx.android, "adb_path", "adb")
    device_id = getattr(ctx.android, "device_id", "") or select_adb_device()
    result = subprocess.run([adb_path, "-s", device_id, "shell", "ime", "list", "-s"], capture_output=True, text=True, timeout=10)
    if "com.android.adbkeyboard/.AdbIME" not in result.stdout:
        raise RuntimeError(
            "ADBKeyboard is not installed. Install senzhk/ADBKeyBoard APK and enable "
            "com.android.adbkeyboard/.AdbIME for Unicode search input."
        )
    subprocess.run([adb_path, "-s", device_id, "shell", "ime", "set", ADB_KEYBOARD_IME], capture_output=True, text=True, timeout=10)


def send_unicode_text(ctx: LiveContext, value: str) -> None:
    adb_path = getattr(ctx.android, "adb_path", "adb")
    device_id = getattr(ctx.android, "device_id", "") or select_adb_device()
    encoded = base64.b64encode(value.encode("utf-8")).decode("ascii")
    subprocess.run(
        [adb_path, "-s", device_id, "shell", "am", "broadcast", "-a", "ADB_INPUT_B64", "--es", "msg", encoded],
        capture_output=True,
        text=True,
        timeout=10,
    )


def delete_search_text(ctx: LiveContext, *, repeats: int) -> None:
    adb_path = getattr(ctx.android, "adb_path", "adb")
    device_id = getattr(ctx.android, "device_id", "") or select_adb_device()
    subprocess.run(
        [adb_path, "-s", device_id, "shell", "input", "keyevent", "123"],
        capture_output=True,
        text=True,
        timeout=10,
    )
    for _ in range(repeats):
        subprocess.run(
            [adb_path, "-s", device_id, "shell", "input", "keyevent", "67"],
            capture_output=True,
            text=True,
            timeout=10,
        )


def find_line_center(ocr_result: dict[str, Any], needle: str) -> dict[str, int] | None:
    needle_lower = needle.lower()
    for line in extract_ocr_lines(ocr_result):
        if needle_lower in line.text.lower():
            return {"x": (line.x_min + line.x_max) // 2, "y": (line.y_min + line.y_max) // 2, "text": line.text}
    return None


def find_bottom_line_center(ocr_result: dict[str, Any], needle: str) -> dict[str, int] | None:
    needle_lower = needle.lower()
    matches = [
        line
        for line in extract_ocr_lines(ocr_result)
        if needle_lower in line.text.lower()
    ]
    if not matches:
        return None
    line = max(matches, key=lambda item: item.y_min)
    return {"x": (line.x_min + line.x_max) // 2, "y": (line.y_min + line.y_max) // 2, "text": line.text}


def search_field_y(ocr_result: dict[str, Any]) -> int | None:
    placeholder = find_line_center(ocr_result, "найти игрока")
    if placeholder:
        return placeholder["y"]
    lines = extract_ocr_lines(ocr_result)
    top = find_line_center(ocr_result, "только мои напарники")
    bottom = find_line_center(ocr_result, "доступно для приглашения")
    if not top or not bottom:
        return None
    candidates = []
    for line in lines:
        center_y = (line.y_min + line.y_max) // 2
        lower = line.text.lower().strip()
        if not (top["y"] < center_y < bottom["y"]):
            continue
        if "сбросить" in lower or "только мои" in lower or "доступно" in lower:
            continue
        if re.search(r"[A-Za-zА-Яа-яЁё]{2,}", line.text):
            candidates.append(line)
    if not candidates:
        return None
    line = min(candidates, key=lambda item: item.y_min)
    return (line.y_min + line.y_max) // 2


def list_active_my_tournaments(conn, *, limit: int = 20):
    init_invite_db(conn)
    return conn.execute(
        """
        SELECT *
        FROM my_tournaments
        WHERE source_status = 'active'
          AND (starts_at IS NULL OR starts_at = '' OR starts_at >= ?)
        ORDER BY starts_at, location
        LIMIT ?
        """,
        (now_iso(), limit),
    ).fetchall()


def now_iso() -> str:
    return datetime.now(MSK).isoformat(timespec="seconds")


if __name__ == "__main__":
    raise SystemExit(main())
